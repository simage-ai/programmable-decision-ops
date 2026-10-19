"""Document provenance: indexing reports into sections, locators, the locators_current condition, and revisions."""

from __future__ import annotations

import json
import shutil

import pytest
from conftest import VEHICLE
from typer.testing import CliRunner

from pdo import tools
from pdo.canonical import object_digest
from pdo.cli import app
from pdo.docs import DocumentError, QuoteError, compare, index_document, locate, normalize_text, reanchor
from pdo.episodes import Program
from pdo.graph import DECISION_KINDS, impact_set, state_digest
from pdo.load import check_state, index, load_state
from pdo.provenance import states_value, verify

DOCS = VEHICLE / "docs"
REPORT, REPORT2, REPORT3 = DOCS / "reliability-test-report.docx", DOCS / "reliability-test-report-rev2.docx", \
    DOCS / "reliability-test-report-rev3.docx"
SHEETS = DOCS / "configuration-sheets.pdf"


@pytest.fixture(scope="module")
def program():
    return Program.load(VEHICLE)


@pytest.fixture(scope="module")
def evaluated(program):
    return {e.episode: e for e in program.evaluate_all(mode="cutoff", verify=True)}


# --- indexing -----------------------------------------------------------------------------------------

def test_word_pdf_and_markdown_give_the_same_sections():
    for built, source in ((REPORT, DOCS / "src/reliability-test-report-rev1.md"),
                          (SHEETS, DOCS / "src/configuration-sheets.md")):
        a, b = index_document(built), index_document(source)
        assert [(s.id, s.path, s.digest) for s in a.sections] == [(s.id, s.path, s.digest) for s in b.sections]
        assert a.digest == b.digest


def test_sections_have_outline_ids_heading_paths_and_pages():
    report = index_document(REPORT)
    assert report.title == "Reliability growth test report, phase 1 (synthetic)"
    assert [s.id for s in report.sections] == ["1", "2", "3", "3.1", "3.2", "3.3", "3.4", "4"]
    assert report.section("3.2").path_text == "3 Results > 3.2 Vehicle B"
    assert report.section("Vehicle B") is report.section("3 Results > 3.2 Vehicle B")
    assert report.section("3.4").paragraphs[1] == "A | 2,000 | 3 | 0.97"  # Word tables become rows
    sheets = index_document(SHEETS)
    assert sheets.pages == 1 and sheets.section("3").page_range == (1, 1)
    assert all("1" != p for s in sheets.sections for p in s.paragraphs)  # the page number is dropped
    assert not [s for s in sheets.sections if s.id == "0"]  # so is the printed title


def test_a_pdf_without_bookmarks_falls_back_to_numbered_headings(tmp_path):
    from pypdf import PdfReader, PdfWriter

    writer = PdfWriter()
    for page in PdfReader(str(SHEETS)).pages:
        writer.add_page(page)
    bare = tmp_path / "sheets.pdf"
    with bare.open("wb") as f:
        writer.write(f)
    # without the metadata title, the printed title is text before the first heading (section 0)
    assert [(s.id, s.digest) for s in index_document(bare).sections if s.id != "0"] == \
        [(s.id, s.digest) for s in index_document(SHEETS).sections]


def test_text_normalization_folds_typography():
    assert normalize_text("“It’s” — 0.95  \n threshold") == "\"It's\" - 0.95 threshold"


def test_unreadable_documents(tmp_path):
    with pytest.raises(DocumentError, match="no document"):
        index_document(tmp_path / "missing.docx")
    (tmp_path / "x.rtf").write_text("x")
    with pytest.raises(DocumentError, match="unsupported"):
        index_document(tmp_path / "x.rtf")
    (tmp_path / "bad.docx").write_text("not a zip")
    with pytest.raises(DocumentError, match="not a readable Word document"):
        index_document(tmp_path / "bad.docx")


# --- locating quotes ----------------------------------------------------------------------------------

def test_locate_returns_a_text_quote_and_position():
    report = index_document(REPORT)
    loc = locate(report, "Its demonstrated mission reliability is 0.98.")
    assert loc["section"] == "3.2" and loc["section_digest"] == report.section("3.2").digest
    assert loc["document_digest"] == report.digest and loc["paragraph"] == 1
    text = report.section("3.2").text
    assert text[loc["char_start"]:loc["char_end"]] == loc["quote"]
    assert text.endswith(loc["quote_prefix"] + loc["quote"] + loc["quote_suffix"])
    sheets = index_document(SHEETS)
    assert locate(sheets, "usable fuel capacity of 45 gal")["page"] == 1


def test_fabricated_ambiguous_and_short_quotes_are_refused():
    report = index_document(REPORT)
    with pytest.raises(QuoteError, match="not in") as info:
        locate(report, "Its demonstrated mission reliability is 0.99.")
    assert "section 3.2" in info.value.hint or "closest" in info.value.hint
    with pytest.raises(QuoteError, match="occurs 3 times"):
        locate(report, "Its demonstrated mission reliability is")
    assert locate(report, "Its demonstrated mission reliability is", section="3.1")["section"] == "3.1"
    with pytest.raises(QuoteError, match="too short"):
        locate(report, "0.9")


def test_reanchor_moves_only_unchanged_sections():
    v1, v2 = index_document(REPORT), index_document(REPORT2)
    a = locate(v1, "Its demonstrated mission reliability is 0.97.")
    moved, info = reanchor(a, v2)
    assert info["status"] == "unchanged" and moved["document_digest"] == v2.digest
    c = locate(v1, "Its demonstrated mission reliability is 0.94, below the 0.95 threshold.")
    assert reanchor(c, v2) == (None, {"status": "changed", "section": "3.3", "quote_found_in": [],
                                      "quote_in_section": False})
    assert {r["path"]: r["status"] for r in compare(v1, v2)} == {
        "1 Purpose": "unchanged", "2 Test conditions": "changed", "3 Results > 3.1 Vehicle A": "unchanged",
        "3 Results > 3.2 Vehicle B": "unchanged", "3 Results > 3.3 Vehicle C": "changed",
        "3 Results > 3.4 Summary": "unchanged", "4 Limitations": "unchanged"}


def test_renamed_and_renumbered_sections_keep_their_identity(tmp_path):
    (tmp_path / "v1.md").write_text("# 1 Intro\n\nHello there.\n\n# 2 Data\n\nThe range is 320 mi.\n")
    (tmp_path / "v2.md").write_text("# 1 Intro\n\nHello there.\n\n# 2 New section\n\nInserted text.\n\n"
                                    "# 3 Data and results\n\nThe range is 320 mi.\n")
    v1, v2 = index_document(tmp_path / "v1.md"), index_document(tmp_path / "v2.md")
    loc = locate(v1, "The range is 320 mi.")
    # renamed and renumbered: nothing identifies "3 Data and results" as "2 Data", so it goes to review
    assert reanchor(loc, v2, v1) == (None, {"status": "ambiguous", "section": None, "quote_found_in": ["3"],
                                            "quote_in_section": False})
    (tmp_path / "v3.md").write_text("# 1 Intro\n\nHello there.\n\n# 2 New section\n\nInserted text.\n\n"
                                    "# 3 Data\n\nThe range is 320 mi.\n")
    v3 = index_document(tmp_path / "v3.md")
    new, info = reanchor(loc, v3, v1)  # renumbered, same heading: moved
    assert info["status"] == "moved" and new["section"] == "3"
    (tmp_path / "v4.md").write_text("# 1 Intro\n\nHello there.\n\n# 2 Data and results\n\nThe range is 320 mi.\n")
    new, info = reanchor(loc, index_document(tmp_path / "v4.md"), v1)  # renamed, same number and text: moved
    assert info["status"] == "moved" and new["section_path"] == "2 Data and results"
    assert {r["status"] for r in compare(v1, v3)} == {"unchanged", "moved", "added"}


def test_the_value_must_appear_in_the_quote():
    assert states_value("Its demonstrated mission reliability is 0.97.", 0.97)
    assert states_value("2,000 miles", 2000) and states_value("97% of missions", 0.97)
    assert not states_value("Its demonstrated mission reliability is 0.97.", 0.98)


# --- the kernel: locators are provenance, gated by readiness -------------------------------------------

def test_locators_are_unhashable(program):
    state = program.materialize("ep:e00")
    obs = index(state)["obs:C.reliability"]
    moved = obs.model_copy(update={"locator": obs.locator.model_copy(update={"section": "9.9"})})
    assert object_digest(obs) == object_digest(moved)
    assert state_digest(program.materialize("ep:e08")) == state_digest(program.materialize("ep:r1")) \
        == state_digest(program.materialize("ep:r2"))


def test_a_revision_flags_only_the_changed_section_and_invalidates_nothing(program, evaluated):
    r1 = evaluated["ep:r1"].run
    cond = r1.value("cond/rc:locators-current")
    assert cond["status"] == "unknown" and "obs:C.reliability (section 3.3" in cond["reason"]
    assert [a["target"] for a in cond["actions"]] == ["obs:C.reliability"]
    assert r1.readiness["state"] == "pending" and r1.actionable["actionable"] is False
    e08 = evaluated["ep:e08"].run
    assert e08.actionable["actionable"] is True
    # no measure, run, check, value, utility, feasibility or decision node is touched; besides the condition,
    # the impact set holds only the clock (the episode's date) and what reads it
    impact = impact_set(e08.graph, r1.graph)
    assert not {n for n in impact if r1.graph.nodes[n].kind in DECISION_KINDS}
    assert impact == {"cond/rc:locators-current", "clock", "cond/rc:evidence-current", "readiness", "actionable"}
    from pdo.actions import pending_actions

    review = [a for a in pending_actions(r1) if a.kind == "review_locator"]
    assert [(a.target, a.role) for a in review] == [("obs:C.reliability", "role:test-engineer")]


def test_re_anchoring_restores_readiness_and_the_approval_holds(evaluated):
    r2 = evaluated["ep:r2"].run
    assert r2.readiness["state"] == "ready" and r2.actionable["actionable"] is True
    r4 = evaluated["ep:r4"].run
    assert r4.decision["recommended"] == ["alt:C"] and r4.actionable["actionable"] is False
    assert all(e.verified for e in evaluated.values() if e.parent)


def test_rejected_and_unindexed_observations_are_ignored(program):
    from pdo.graph import compile_graph
    from pdo.run import run

    state = program.materialize("ep:r1").model_copy(deep=True)  # never mutate the program's cached states
    objects = index(state)
    objects["obs:C.reliability"].applicability = "rejected"  # a stale locator on a rejected value
    assert run(compile_graph(state, program.base_dir)).value("cond/rc:locators-current")["status"] == "pass"
    state = program.materialize("ep:r1").model_copy(deep=True)
    index(state)["ev:reliability-C"].document_digest = None  # no longer an indexed document
    issues = [str(i) for i in check_state(state)]
    assert any("is not an indexed document" in i and i.startswith("warning") for i in issues)
    assert run(compile_graph(state, program.base_dir)).value("cond/rc:locators-current")["status"] == "pass"


# --- verification ---------------------------------------------------------------------------------------

def copy_vehicle(tmp_path):
    dest = tmp_path / "vehicle"
    shutil.copytree(VEHICLE, dest, ignore=shutil.ignore_patterns("__pycache__"))
    return dest


def test_verify_accepts_the_example_and_catches_a_doctored_locator(tmp_path):
    assert verify(load_state(VEHICLE).state, VEHICLE) == []
    dest = copy_vehicle(tmp_path)
    f = dest / "baseline" / "evidence.yaml"
    f.write_text(f.read_text().replace("quote: Its demonstrated mission reliability is 0.97.",
                                       "quote: Its demonstrated mission reliability is 0.99."))
    problems = verify(load_state(dest).state, dest)
    assert [(p["severity"], p["where"]) for p in problems] == [("error", "obs:A.reliability.locator"),
                                                               ("warning", "obs:A.reliability.locator")]
    assert "not in section 3.1" in problems[0]["message"] and "does not state the value 0.97" in problems[1]["message"]
    r = tools.validate(dest)
    assert r["valid"] is False and any("obs:A.reliability.locator" in e for e in r["errors"])


def test_verify_warns_when_a_value_drifts_from_its_quote_or_the_file_changes(tmp_path):
    dest = copy_vehicle(tmp_path)
    f = dest / "baseline" / "evidence.yaml"
    f.write_text(f.read_text().replace("meas:reliability, value: 0.97,", "meas:reliability, value: 0.975,"))
    problems = verify(load_state(dest).state, dest)
    assert [p["severity"] for p in problems] == ["warning"] and "does not state the value" in problems[0]["message"]
    shutil.copy(REPORT2, dest / "docs" / "reliability-test-report.docx")  # a new version over the old file
    messages = [p["message"] for p in verify(load_state(dest).state, dest)]
    assert any("not the version indexed" in m for m in messages)


# --- tools, CLI and MCP --------------------------------------------------------------------------------

def test_docs_tools_index_read_and_locate():
    r = tools.docs_index(SHEETS)
    assert r["ok"] and r["format"] == "pdf" and [s["id"] for s in r["sections"]] == ["1", "2", "3", "4", "5"]
    r = tools.docs_read(REPORT, "3.3")
    assert r["sections"][0]["text"][0].startswith("Vehicle C, the low-cost commercial platform")
    r = tools.docs_locate(REPORT, "mission reliability is 0.98")
    assert r["locator"]["section"] == "3.2" and r["numbers"] == [0.98]
    bad = tools.docs_locate(REPORT, "mission reliability is 0.99")
    assert bad["ok"] is False and "not in" in bad["error"] and "pdo docs show" in bad["hint"]


def test_extract_adds_provisional_located_values_and_refuses_fabrication(tmp_path):
    dest = copy_vehicle(tmp_path)
    proposals = {"evidence": "ev:reliability-A", "observations": [
        {"id": "obs:A.miles", "alternative": "alt:A", "measure": "meas:reliability", "value": 0.97,
         "quote": "Its demonstrated mission reliability is 0.97."}]}
    r = tools.docs_extract(dest, proposals, parent="ep:e00")
    assert r["ok"], r
    obs = r["change"]["add"]["observations"][0]
    assert obs["applicability"] == "provisional" and obs["locator"]["section"] == "3.1"
    assert obs["locator"]["document_path"] == "docs/reliability-test-report.docx"
    fake = {"evidence": "ev:reliability-A", "observations": [
        {"id": "obs:A.reliability", "value": 0.99, "quote": "Its demonstrated mission reliability is 0.99."}]}
    before = set((dest / "changes").iterdir())
    r = tools.docs_extract(dest, fake, parent="ep:e00", write=True)
    assert r["ok"] is False and "not in" in r["error"] and set((dest / "changes").iterdir()) == before


def test_revise_writes_an_episode_that_moves_unchanged_locators(tmp_path):
    dest = copy_vehicle(tmp_path)
    shutil.copy(REPORT3, dest / "docs" / "report-rev3.docx")
    r = tools.docs_revise(dest, "ev:reliability-A", "docs/report-rev3.docx", parent="ep:e00",
                          role="role:test-engineer", write=True)
    assert r["ok"], r
    rev = r["revision"]
    assert {x["observation"] for x in rev["reanchored"]} == {"obs:A.reliability", "obs:B.reliability"}
    assert [x["observation"] for x in rev["flagged"]] == ["obs:C.reliability"]
    assert r["approval_invalidated"] is False and r["after"]["readiness"] == "pending"
    assert r["pending_actions_after"][0]["resolve_with"]["tool"] == "pdo_docs_extract"
    reloaded = Program.load(dest)
    assert r["episode"] in reloaded.changes
    again = tools.docs_revise(dest, "ev:reliability-A", "docs/report-rev3.docx", parent=r["episode"])
    assert again["ok"] is False and "is the file being revised" in again["error"]


def test_cli_docs_commands():
    runner = CliRunner()
    out = runner.invoke(app, ["docs", "index", str(REPORT)])
    assert out.exit_code == 0 and "3 Results > 3.2 Vehicle B" in out.output
    out = runner.invoke(app, ["docs", "show", str(REPORT), "3.3"])
    assert out.exit_code == 0 and "below the 0.95 threshold" in out.output
    out = runner.invoke(app, ["docs", "locate", str(SHEETS), "usable fuel capacity of 47.5 gal", "--json"])
    assert out.exit_code == 0 and json.loads(out.output)["locator"]["section"] == "4"
    out = runner.invoke(app, ["docs", "locate", str(SHEETS), "usable fuel capacity of 99 gal"])
    assert out.exit_code == 1
    out = runner.invoke(app, ["docs", "verify", str(VEHICLE)])
    assert out.exit_code == 0 and "0 error(s)" in out.output
    out = runner.invoke(app, ["docs", "revise", str(VEHICLE), "ev:reliability-C", str(REPORT3), "--parent", "ep:e00"])
    assert out.exit_code == 0 and "REVIEW    obs:C.reliability" in out.output and "preview only" in out.output


def test_explain_package_and_dashboard_show_the_passage(program, evaluated):
    from pdo.package import episode_record
    from pdo.report import explain_node

    r1 = evaluated["ep:r1"].run
    src = explain_node(r1, "y/alt:C/meas:reliability")["sources"][0]
    assert src["document"] == "docs/reliability-test-report.docx" and src["current"] is False
    assert explain_node(r1, "y/alt:A/meas:reliability")["sources"][0]["current"] is True
    record = episode_record(r1, episode="ep:r1")
    claim = next(c for c in record.claims if c.produced_by == "obs:C.reliability")
    assert "§ 3.3" in claim.source_locator and "awaiting review" in claim.source_locator
    from pdo.dashboard import _node_state

    node = _node_state(r1.graph, r1, "y/alt:C/meas:reliability", {})
    assert node["src"][0]["section"] == "3.3" and node["src"][0]["current"] is False


# --- the provenance benchmark ----------------------------------------------------------------------------

def test_provenance_benchmark_passes_its_gate():
    from conftest import ROOT

    from pdo.bench.provenance import run_provenance

    result = run_provenance(ROOT / "bench" / "provenance")
    t = result.totals()
    n = t["cases"]
    assert result.passed and t["agree"] == (n, n) and t["changed_detected"] == (n, n)
    assert t["flagged"] < t["whole_document"] and t["changed_flagged"] == t["changed_passages"]
    assert t["approvals_kept"] == (n, n) and t["no_decision_nodes"] == (n, n)
    twins = {r["case"]: r["flagged"] for r in result.rows}  # old text surviving elsewhere never hides a change
    assert twins["rel-c-superseded-appendix"] == ["obs:C.reliability"]
    assert twins["rel-a-text-moves-to-b"] == ["obs:A.reliability", "obs:B.reliability"]


def test_benchmark_edits_and_word_writer(tmp_path):
    from pdo.bench.provenance import apply_edits, write_docx

    source = (DOCS / "src" / "configuration-sheets.md").read_text()
    moved = apply_edits(source, [{"move": "# 4 Vehicle C: low-cost commercial platform",
                                  "before": "# 2 Vehicle A: current production derivative"}])
    assert moved.index("# 4 Vehicle C") < moved.index("# 2 Vehicle A") and len(moved) == len(source)
    with pytest.raises(ValueError, match="matches 0 times"):
        apply_edits(source, [{"replace": "not in the text", "with": "x"}])
    write_docx(source, tmp_path / "sheets.docx")
    assert index_document(tmp_path / "sheets.docx").digest == index_document(SHEETS).digest


def test_authored_lines_leave_out_what_tools_wrote(tmp_path):
    from pdo.bench.suite import suite_yaml_lines

    (tmp_path / "baseline").mkdir()
    (tmp_path / "changes").mkdir()
    (tmp_path / "program.yaml").write_text("id: prog:x\n")
    (tmp_path / "baseline" / "a.yaml").write_text(
        "observations:\n"
        "  - {id: obs:x, value: 1,\n"
        "     locator: {document_digest: d, section: '1',\n"
        "       quote: q}}\n"
        "  - id: obs:y\n"
        "    locator:\n"
        "      section: '2'\n"
        "      quote: q\n"
        "    value: 2\n")
    (tmp_path / "changes" / "c.yaml").write_text("# Written by pdo docs revise\nid: ep:c\nparent: ep:b\n")
    assert suite_yaml_lines(tmp_path) == 5  # program.yaml, observations:, obs:x, obs:y, value: 2


def test_skill_is_well_formed_and_names_real_commands():
    import re

    import yaml
    from conftest import ROOT

    from pdo.mcp_server import READ_ONLY, WRITES

    text = (ROOT / "skills" / "document-provenance" / "SKILL.md").read_text()
    front = yaml.safe_load(text.split("---")[1])
    assert front["name"] == "document-provenance" and "provenance" in front["description"]
    commands = {c.name for c in app.registered_groups if c.name == "docs"}
    assert commands == {"docs"}
    for sub in set(re.findall(r"`pdo docs (\w+)", text)):
        assert CliRunner().invoke(app, ["docs", sub, "--help"]).exit_code == 0, sub
    tools_named = set(re.findall(r"`(pdo_\w+)`", text))
    assert tools_named <= {fn.__name__ for fn in (*READ_ONLY, *WRITES)}, tools_named
    assert "locators_current" in text and "applicability: provisional" in text


# --- the verifier's findings, each pinned by a test ------------------------------------------------------

def md(tmp_path, name, text):
    f = tmp_path / name
    f.write_text(text)
    return index_document(f)


def test_a_changed_section_never_follows_its_old_text_elsewhere(tmp_path):
    v1 = md(tmp_path, "v1.md", "# 1 Vehicle A\n\nMission reliability: 0.97.\n\n# 2 Vehicle D (reference design)\n\n"
                               "Mission reliability: 0.97.\n")
    v2 = md(tmp_path, "v2.md", "# 1 Vehicle A\n\nMission reliability: 0.91.\n\n# 2 Vehicle D (reference design)\n\n"
                               "Mission reliability: 0.97.\n")
    loc = locate(v1, "Mission reliability: 0.97.", section="1")
    new, info = reanchor(loc, v2, v1)
    assert new is None and info["status"] == "changed" and info["section"] == "1"
    assert {r["path"]: r["status"] for r in compare(v1, v2)} == {"1 Vehicle A": "changed",
                                                                 "2 Vehicle D (reference design)": "unchanged"}
    # renamed under the same number with the same text: moved; renamed and edited, with its old text
    # surviving in an appendix: ambiguous, so it goes to review instead of following the old text
    v3 = md(tmp_path, "v3.md", "# 1 Vehicle A (production)\n\nMission reliability: 0.97.\n\n# 2 Vehicle D\n\n"
                               "Mission reliability: 0.97.\n")
    assert reanchor(loc, v3, v1)[1] == {"status": "moved", "section": "1"}
    v4 = md(tmp_path, "v4.md", "# 1 Vehicle A (reclassified)\n\nMission reliability: 0.96.\n\n"
                               "# 2 Vehicle D (reference design)\n\nMission reliability: 0.97.\n\n"
                               "# Appendix\n\nOld text. Mission reliability: 0.97.\n")
    assert reanchor(loc, v4, v1)[0] is None


def test_quotes_must_be_whole_words_and_numbers():
    from pdo.docs import _bounded
    from pdo.provenance import numbers_in

    report = index_document(REPORT)
    with pytest.raises(QuoteError, match="whole words and numbers"):
        locate(report, "mission reliability is 0.9", section="3.2")
    with pytest.raises(QuoteError):
        locate(index_document(SHEETS), "usable fuel capacity of 4", section="2")
    assert locate(report, "mission reliability is 0.98", section="3.2")["quote"] == "mission reliability is 0.98"
    assert not _bounded("is 0.98.", 0, 6) and _bounded("is 0.98.", 0, 7) and _bounded("is 0.98.", 0, 8)
    assert not _bounded("1,000 units", 0, 2) and not _bounded("phase-2 test", 6, 12)
    assert numbers_in("between 5-38 °C, -4 at night, 2,000 miles") == [5, 38, -4, 2000]


def test_extract_refuses_a_contradicting_quote_and_an_unaccepted_re_anchor(tmp_path, program):
    dest = copy_vehicle(tmp_path)
    contradict = {"evidence": "ev:reliability-C", "observations": [
        {"id": "obs:C.reliability", "value": 0.94, "quote": "Its demonstrated mission reliability is 0.96, above "
                                                            "the 0.95 threshold."}]}
    r = tools.docs_extract(dest, contradict, parent="ep:r3")
    assert r["ok"] is False and "does not state 0.94" in r["error"]
    stale = {"observations": [{"id": "obs:C.reliability", "value": 0.96,
                               "quote": "Its demonstrated mission reliability is 0.96, above the 0.95 threshold."}]}
    r = tools.docs_extract(dest, stale, parent="ep:r3")  # evidence defaults to the observation's own
    assert r["ok"] is False and "evidence owner's judgement" in r["error"]
    r = tools.docs_extract(dest, {**stale, "accept": True}, parent="ep:r3")
    assert r["ok"] is False and "needs the role" in r["error"]
    r = tools.docs_extract(dest, {**stale, "accept": True}, parent="ep:r3", role="role:test-engineer")
    assert r["ok"] and r["rows"][0]["action"] == "update" and r["after"]["recommended"] == ["alt:C"]


def test_accept_accepts_and_new_documents_stay_inside_the_directory(tmp_path):
    dest = copy_vehicle(tmp_path)
    accept = {"accept": True, "observations": [{"id": "obs:B.reliability", "value": 0.98,
                                                "quote": "Its demonstrated mission reliability is 0.98."}]}
    r = tools.docs_extract(dest, accept, parent="ep:e04", role="role:test-engineer")  # rejected in e04
    assert r["ok"] and {"target": "obs:B.reliability", "slot": "applicability", "new_value": "accepted"} in \
        [a.model_dump() if hasattr(a, "model_dump") else a for a in r["change"]["assign"]]
    outside = {"evidence_items": [{"id": "ev:readme", "document_path": "../../README.md"}],
               "observations": [{"id": "obs:x", "alternative": "alt:A", "measure": "meas:reliability",
                                 "value": 1, "evidence": "ev:readme", "quote": "anything at all"}]}
    r = tools.docs_extract(dest, outside, parent="ep:e00")
    assert r["ok"] is False and "outside" in r["error"]


def test_a_value_from_an_indexed_document_must_be_located(tmp_path, program):
    from pdo.graph import compile_graph
    from pdo.run import run

    state = program.materialize("ep:r3").model_copy(deep=True)
    index(state)["obs:C.reliability"].locator = None  # deleting the locator does not clear the gate
    result = run(compile_graph(state, program.base_dir))
    cond = result.value("cond/rc:locators-current")
    assert cond["status"] == "unknown" and "without a locator: obs:C.reliability" in cond["reason"]
    assert any(p["where"] == "obs:C.reliability.locator" for p in verify(state, program.base_dir))
    dest = copy_vehicle(tmp_path)
    f = dest / "baseline" / "evidence.yaml"
    text = f.read_text()
    start = text.index("     locator: {document_path: docs/reliability-test-report.docx")
    end = text.index("}}", start) + 2
    f.write_text(text[:start].rstrip().rstrip(",") + "}" + text[end:])  # obs:A.reliability without its locator
    shutil.copy(REPORT2, dest / "docs" / "r2.docx")
    r = tools.docs_revise(dest, "ev:reliability-A", "docs/r2.docx", parent="ep:e00")
    assert r["ok"] and {"observation": "obs:A.reliability", "status": "not located"}.items() <= \
        next(x for x in r["revision"]["flagged"] if x["observation"] == "obs:A.reliability").items()


def test_verify_errors_on_a_replaced_file_or_doctored_fields(tmp_path):
    dest = copy_vehicle(tmp_path)
    shutil.copy(REPORT2, dest / "docs" / "reliability-test-report.docx")  # a new version copied over the old
    problems = verify(load_state(dest).state, dest)
    assert {p["severity"] for p in problems if "reliability" in p["where"]} == {"error"}
    dest2 = copy_vehicle(tmp_path / "b")
    f = dest2 / "baseline" / "evidence.yaml"
    text = f.read_text().replace("section_path: 3 Results > 3.1 Vehicle A", "section_path: 3 Results > 3.2 Vehicle B")
    f.write_text(text)
    problems = verify(load_state(dest2).state, dest2)
    assert any(p["severity"] == "error" and "section_path" in p["message"] for p in problems)


def test_cli_validate_checks_documents(tmp_path):
    dest = copy_vehicle(tmp_path)
    f = dest / "baseline" / "evidence.yaml"
    f.write_text(f.read_text().replace("quote: Its demonstrated mission reliability is 0.98.",
                                       "quote: Its demonstrated mission reliability is 0.999."))
    out = CliRunner().invoke(app, ["validate", str(dest)])
    assert out.exit_code == 1 and "obs:B.reliability.locator" in out.output


def test_package_reports_document_problems():
    r = tools.package(VEHICLE, "ep:r1")
    assert r["ok"] and r["document_problems"] == []


def test_section_ids_are_printed_numbers(tmp_path):
    ix = md(tmp_path, "a.md", "# Executive summary\n\nIntro.\n\n# 1 Vehicle A\n\nSame sentence.\n\n"
                              "# 2 Vehicle B\n\nOther sentence.\n")
    assert [s.id for s in ix.sections] == ["s1", "1", "2"]
    assert ix.section("2").heading == "2 Vehicle B" and ix.section("Vehicle B").id == "2"
    plain = md(tmp_path, "b.md", "# Intro\n\nx y.\n\n## Detail\n\nz w.\n")
    assert [s.id for s in plain.sections] == ["1", "1.1"]


def test_word_tables_keep_repeated_cells(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("1 Data", level=1)
    t = d.add_table(rows=1, cols=3)
    for j, text in enumerate(["A", "2,000", "2,000"]):
        t.cell(0, j).text = text
    d.save(str(tmp_path / "t.docx"))
    assert index_document(tmp_path / "t.docx").sections[0].paragraphs == ["A | 2,000 | 2,000"]


def test_normalization_keeps_exponents_and_ignores_line_break_hyphens(tmp_path):
    assert normalize_text("MTBF of 10⁵ h; 0.97¹") == "MTBF of 10^5 h; 0.97^1"
    a = md(tmp_path, "a.md", "# 1 Course\n\nThe mixed-terrain endurance course.\n")
    b = md(tmp_path, "b.txt", "1 Course\n\nThe mixedterrain endurance course.\n")
    assert a.sections[0].digest == b.sections[0].digest
    assert locate(b, "mixed-terrain endurance")["quote"] == "mixedterrain endurance"


def test_markdown_inline_markup_setext_fences_and_comments(tmp_path):
    ix = md(tmp_path, "m.md", "Results\n=======\n\nReliability is **0.97**, see [Table 2](#t2) and `MR`.\n"
                              "<!-- a comment -->\n\n~~~\n# not a heading\n~~~\n\nDetails\n-------\n\nMore.\n")
    assert [(s.id, s.heading) for s in ix.sections] == [("1", "Results"), ("1.1", "Details")]
    assert ix.sections[0].paragraphs == ["Reliability is 0.97, see Table 2 and MR.", "# not a heading"]


def test_numbered_headings_without_an_outline_form_a_chain():
    from pdo.docs import _numbered_headings

    lines = list(enumerate(["1 Purpose", "Text.", "2 Results", "2.1 Vehicle A", "Rows:", "2 Suspension",
                            "4 Driveline", "2.2 Vehicle B", "3 Limitations", "Text."]))
    assert sorted(v[1] for v in _numbered_headings(lines).values()) == \
        ["1 Purpose", "2 Results", "2.1 Vehicle A", "2.2 Vehicle B", "3 Limitations"]


def test_running_headers_of_a_two_page_pdf_are_dropped():
    from pdo.docs import _strip_furniture

    pages = [["Reliability report, revision 1", "1 Purpose", "Text runs"], ["Reliability report, revision 1",
                                                                           "on here.", "2"]]
    assert _strip_furniture(pages, None) == [["1 Purpose", "Text runs"], ["on here."]]
    # lines that differ only in a number that does not track the page are content, not furniture
    content = [["1 Results", "Mission reliability: 0.97"], ["2 More", "Mission reliability: 0.98"]]
    assert _strip_furniture(content, None) == content
    # nor is a line that prints an outline heading, even when it starts every page
    pages = [["Configuration 1", "Usable fuel capacity: 40 gal"], ["Configuration 2", "Usable fuel capacity: 45 gal"],
             ["Configuration 3", "Usable fuel capacity: 47 gal"]]
    assert _strip_furniture(pages, None, ["Configuration 1", "Configuration 2", "Configuration 3"]) == pages
    assert _strip_furniture([["Page 1 of 3", "a"], ["Page 2 of 3", "b"], ["Page 3 of 3", "c"]], None) == \
        [["a"], ["b"], ["c"]]


def test_empty_headings_and_the_pdf_title(tmp_path):
    import docx

    d = docx.Document()
    d.add_heading("1 Results", level=1)
    d.add_paragraph("First part.")
    d.add_heading("", level=2)
    d.add_paragraph("Second part.")
    d.save(str(tmp_path / "e.docx"))
    assert [(s.id, s.paragraphs) for s in index_document(tmp_path / "e.docx").sections] == \
        [("1", ["First part.", "Second part."])]
    assert index_document(SHEETS).title == "Vehicle configuration sheets, revision 1 (synthetic)"


def test_a_missing_docs_extra_is_one_warning(monkeypatch):
    from pdo import docs, provenance
    from pdo.docs import DocumentError

    def missing(path):
        raise DocumentError("reading Word documents needs python-docx: uv sync --extra docs")

    monkeypatch.setattr(docs, "_docx", missing)
    monkeypatch.setattr(docs, "_pdf", missing)
    provenance._cached_index.cache_clear()
    try:
        r = tools.validate(VEHICLE)
    finally:
        provenance._cached_index.cache_clear()
    assert r["valid"] and [w for w in r["warnings"] if "docs extra" in w] == \
        ["warning: documents: Word and PDF documents were not checked: install the docs extra (uv sync --extra docs)"]
    assert not [w for w in r["warnings"] if w.startswith("error")]


# --- the second round of findings ----------------------------------------------------------------------

def test_scientific_notation_is_one_number_and_never_cut():
    from pdo.docs import _bounded
    from pdo.provenance import numbers_in, states_value

    assert states_value("The MTBF is 10^5 h.", 100000) and states_value("a rate of 1.2×10^-3 per hour", 0.0012)
    assert states_value("a rate of 1.2e-3 per hour", 0.0012) and not states_value("The MTBF is 10^5 h.", 10)
    assert numbers_in("10^5 h and 3 x 10^2") == [100000, 300]
    text = "The MTBF is 10^5 h; the rate is 1.2×10^-3 or 1.2e-3."
    assert not _bounded(text, 0, text.index("10^") + 2)
    assert not _bounded(text, 0, text.index("1.2×") + 3) and not _bounded(text, 0, text.index("1.2e") + 3)
    assert normalize_text("tank holds 1½ gal") == "tank holds 1 1/2 gal"


def test_pdf_line_end_hyphens_that_belong_to_the_text_are_kept():
    from pdo.docs import _dehyphenate

    assert _dehyphenate(["tested to MIL-", "STD-810 and a 2,000-", "mile course, reliabil-", "ity grew"]) == \
        "tested to MIL-STD-810 and a 2,000-mile course, reliability grew"


def test_accepting_needs_the_evidence_owner_and_re_anchoring_needs_acceptance(tmp_path):
    dest = copy_vehicle(tmp_path)
    fix = {"accept": True, "observations": [{"id": "obs:C.reliability", "value": 0.96,
                                             "quote": "Its demonstrated mission reliability is 0.96, above the 0.95 "
                                                      "threshold."}]}
    r = tools.docs_extract(dest, fix, parent="ep:r3", role="role:nobody")
    assert r["ok"] is False and "not a role" in r["error"]
    r = tools.docs_extract(dest, fix, parent="ep:r3", role="role:cost-analyst")
    assert r["ok"] is False and "only role:test-engineer" in r["error"]
    r = tools.docs_extract(dest, fix, parent="ep:r3", role="role:test-engineer")
    assert r["ok"] and any(a["slot"] == "applicability_rationale" and "Accepted by role:test-engineer" in
                           a["new_value"] for a in r["change"]["assign"])
    # removing the locator first does not open a way around the review
    r = tools.apply_change(dest, {"parent": "ep:r3", "assign": [{"target": "obs:C.reliability", "slot": "locator",
                                                                 "new_value": None}]})
    assert r["ok"] is False and "never typed" in r["error"]


def test_unanchored_values_need_acceptance_to_be_located(tmp_path, program):
    from pdo.provenance import extract

    state = program.materialize("ep:r1").model_copy(deep=True)
    index(state)["obs:C.reliability"].locator = None
    proposal = {"observations": [{"id": "obs:C.reliability", "value": 0.94,
                                  "quote": "Until the board rules, its demonstrated mission reliability is 0.94"}]}
    with pytest.raises(ValueError, match="evidence owner's judgement"):
        extract(state, program.base_dir, proposal)
    change, rows = extract(state, program.base_dir, {**proposal, "accept": True}, role="role:test-engineer")
    assert rows[0]["action"] == "re-anchor"


def test_author_roles_must_exist(tmp_path):
    dest = copy_vehicle(tmp_path)
    r = tools.apply_change(dest, {"parent": "ep:e00", "author_role": "role:nobody",
                                  "assign": [{"target": "obs:B.cost", "slot": "value", "new_value": 1.25}]})
    assert r["ok"] is False and "author_role role:nobody is not a role" in r["error"]


def test_explicit_verify_fails_when_nothing_could_be_checked(monkeypatch):
    from pdo import docs, provenance
    from pdo.docs import DocumentError

    def missing(path):
        raise DocumentError("reading Word documents needs python-docx: uv sync --extra docs")

    monkeypatch.setattr(docs, "_docx", missing)
    monkeypatch.setattr(docs, "_pdf", missing)
    provenance._cached_index.cache_clear()
    try:
        r = tools.docs_verify(VEHICLE, "ep:e00")
    finally:
        provenance._cached_index.cache_clear()
    assert r["unchecked"] > 0 and r["errors"] >= r["unchecked"]


def test_packaging_refuses_a_contradicted_locator(tmp_path):
    dest = copy_vehicle(tmp_path)
    f = dest / "baseline" / "evidence.yaml"
    f.write_text(f.read_text().replace("quote: Its demonstrated mission reliability is 0.98.",
                                       "quote: Its demonstrated mission reliability is 0.999."))
    r = tools.package(dest, "ep:e00", out=tmp_path / "pkg")
    assert r["ok"] is False and "obs:B.reliability.locator" in r["error"] and not (tmp_path / "pkg").exists()
    out = CliRunner().invoke(app, ["package", str(dest), "-e", "ep:e00", "-o", str(tmp_path / "pkg2")])
    assert out.exit_code == 1 and not (tmp_path / "pkg2").exists()


def test_revise_by_path_when_the_old_file_is_gone(tmp_path):
    dest = copy_vehicle(tmp_path)
    shutil.copy(REPORT2, dest / "docs" / "r2.docx")
    (dest / "docs" / "reliability-test-report.docx").unlink()
    r = tools.docs_revise(dest, "docs/reliability-test-report.docx", "docs/r2.docx", parent="ep:e00")
    assert r["ok"] and r["revision"]["sections"] is None
    assert [x["observation"] for x in r["revision"]["flagged"]] == ["obs:C.reliability"]
