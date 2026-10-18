"""Provenance benchmark: section-level edits to indexed reports, absorbed with pdo docs revise.

Each case edits a report's Markdown source, writes the new version as a Word file, and revises the program
at the suite's parent episode in a scratch copy. It measures what a person would have to review: the values
flagged (their cited section changed) against the hand labels and against a whole-document policy that
sends back every value citing the revised report; how many flagged values had their quoted passage change
(precision) and whether any changed passage escaped (recall); and that absorbing a revision invalidates
nothing: no approval stops matching and no decision node is in the impact set.
"""

from __future__ import annotations

import re
import shutil
import tempfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import yaml

from pdo import model as m
from pdo.bench.report import tsv, write_if_changed
from pdo.docs import normalize_text
from pdo.graph import DECISION_KINDS, impact_set


def suite_source(suite_file: Path) -> str:
    """The suite's identity for the proposal table: a digest of its cases and labels, and the kernel version."""
    from pdo import __version__
    from pdo.canonical import digest

    return f"suite {suite_file.parent.name} {digest(suite_file.read_text(encoding='utf-8'))}; pdo {__version__}"


@dataclass
class ProvenanceResult:
    suite: str
    parent: str
    rows: list[dict] = field(default_factory=list)
    source: str = ""

    def rate(self, key: str) -> tuple[int, int]:
        return sum(1 for r in self.rows if r[key]), len(self.rows)

    def totals(self) -> dict[str, Any]:
        flagged = sum(len(r["flagged"]) for r in self.rows)
        whole = sum(r["whole_document"] for r in self.rows)
        changed = sum(len(r["expected_changed"]) for r in self.rows)
        caught = sum(len(set(r["expected_changed"]) & set(r["flagged"])) for r in self.rows)
        return {"cases": len(self.rows), "agree": self.rate("agree"), "flagged": flagged, "whole_document": whole,
                "changed_passages": changed, "changed_flagged": caught,
                "changed_detected": self.rate("changed_agree"),
                "approvals_kept": self.rate("approval_kept"), "no_decision_nodes": self.rate("no_decision_nodes"),
                "readiness_pending_iff_flagged": self.rate("readiness_right"), "verified": self.rate("verified")}

    @property
    def passed(self) -> bool:
        t = self.totals()
        return all(t[k][0] == t[k][1] for k in ("agree", "changed_detected", "approvals_kept", "no_decision_nodes",
                                                 "readiness_pending_iff_flagged", "verified")) \
            and t["changed_flagged"] == t["changed_passages"]


def apply_edits(text: str, edits: list[dict]) -> str:
    """Apply replace and move edits in order; each must match exactly once."""
    for e in edits:
        if "replace" in e:
            if text.count(e["replace"]) != 1:
                raise ValueError(f"edit {e['replace'][:40]!r} matches {text.count(e['replace'])} times")
            text = text.replace(e["replace"], e["with"])
        elif "move" in e:
            text = _move(text, e["move"], e["before"])
        else:
            raise ValueError(f"unknown edit {e}")
    return text


def _move(text: str, heading: str, before: str) -> str:
    """Cut the section starting at `heading` (to the next heading of the same or a higher level) and insert
    it before the heading `before`."""
    level = len(heading) - len(heading.lstrip("#"))
    start = text.index(heading + "\n")
    rest = text[start + len(heading):]
    nxt = re.search(r"^#{1," + str(level) + r"}\s", rest, re.M)
    end = start + len(heading) + (nxt.start() if nxt else len(rest))
    block, text = text[start:end], text[:start] + text[end:]
    at = text.index(before + "\n")
    return text[:at] + block + text[at:]


def write_docx(markdown: str, path: Path) -> None:
    """A Word file from the report's Markdown subset: front-matter title, # headings, paragraphs, tables."""
    import docx

    document = docx.Document()
    if markdown.startswith("---\n"):
        end = markdown.find("\n---", 4)
        title = re.search(r"^title:\s*(.+)$", markdown[4:end], re.M)
        if title:
            document.core_properties.title = title.group(1).strip()
            document.add_heading(title.group(1).strip(), level=0)
        markdown = markdown[end + 4:]
    table: list[list[str]] = []

    def flush_table() -> None:
        if table:
            t = document.add_table(rows=len(table), cols=len(table[0]))
            for i, row in enumerate(table):
                for j, cell in enumerate(row):
                    t.cell(i, j).text = cell
            table.clear()

    for block in re.split(r"\n\s*\n", markdown.strip()):
        for line in block.splitlines():
            if line.startswith("|"):
                cells = [c.strip() for c in line.strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                    table.append(cells)
                continue
            flush_table()
            heading = re.match(r"^(#{1,6})\s+(.*)$", line)
            if heading:
                document.add_heading(heading.group(2), level=len(heading.group(1)))
        flush_table()
        body = [ln for ln in block.splitlines() if not ln.startswith(("|", "#"))]
        if body:
            document.add_paragraph(" ".join(ln.strip() for ln in body))
    document.save(str(path))


def run_provenance(suite: str | Path) -> ProvenanceResult:
    from pdo.bench.metrics import approval_invalidated
    from pdo.episodes import Program
    from pdo.provenance import revise

    suite_file = Path(suite)
    if suite_file.is_dir():
        suite_file = suite_file / "cases.yaml"
    spec = yaml.safe_load(suite_file.read_text(encoding="utf-8"))
    source_program = (suite_file.parent / spec["program"]).resolve()
    result = ProvenanceResult(suite_file.parent.name, spec["parent"])
    with tempfile.TemporaryDirectory(prefix="pdo-provenance-") as tmp:
        root = Path(tmp) / "program"
        shutil.copytree(source_program, root, ignore=shutil.ignore_patterns("__pycache__"))
        (root / "docs" / "bench").mkdir(parents=True, exist_ok=True)
        program = Program.load(root)
        parent = spec["parent"]
        before = program.evaluate(parent, verify=False)
        state = program.materialize(parent)
        objects = {o.id: o for o in state.observations}
        evidence = {e.id: e for e in state.evidence_items}
        result.source = suite_source(suite_file)
        for case in spec["cases"]:
            doc = spec["documents"][case["document"]]
            text = apply_edits((root / doc["source"]).read_text(encoding="utf-8"), case.get("edits") or [])
            new_file = root / "docs" / "bench" / f"{case['id']}.docx"
            write_docx(text, new_file)
            citing = sorted(o.id for o in objects.values() if o.locator is not None
                            and (o.applicability or "accepted") != "rejected"
                            and evidence[o.evidence].document_path == doc["path"])
            change_data, report = revise(state, root, doc["path"], new_file)
            change = m.ChangeSet.model_validate({**change_data, "id": "ep:bench", "parent": parent})
            after = program.preview(change, parent=parent, mode="cutoff", verify=True)
            flagged = sorted(f["observation"] for f in report["flagged"] if f["applies"])
            # the system's own view of which quoted passages changed: flagged, and the quote is no longer in the
            # cited section's successor
            detected = sorted(f["observation"] for f in report["flagged"] if f["applies"] and not f["quote_in_section"])
            impact = impact_set(before.run.graph, after.run.graph)
            expect = case["expect"]
            pending = after.run.readiness["state"] == "pending"
            result.rows.append({
                "case": case["id"], "document": case["document"], "family": case["family"],
                "flagged": flagged, "expected": sorted(expect["flagged"]),
                "agree": flagged == sorted(expect["flagged"]),
                "expected_changed": sorted(expect["changed_passages"]), "changed_detected": detected,
                "changed_agree": detected == sorted(expect["changed_passages"]), "whole_document": len(citing),
                "reanchored": len(report["reanchored"]),
                "sections_differing": sum(1 for s in report["sections"] or [] if s["status"] != "unchanged"),
                "readiness": after.run.readiness["state"], "readiness_right": pending == bool(flagged),
                "approval_kept": not approval_invalidated(before.run, after.run),
                "no_decision_nodes": not {n for n in impact if after.run.graph.nodes[n].kind in DECISION_KINDS},
                "verified": bool(after.verified),
            })
    return result


def _pct(num: int, den: int) -> str:
    return f"{100 * num / den:.1f}" if den else ""


def summary_rows(result: ProvenanceResult) -> list[tuple]:
    t = result.totals()
    rows = [("cases", t["cases"], "", ""),
            ("label_agreement_pct", _pct(*t["agree"]), *t["agree"]),
            ("changed_passage_agreement_pct", _pct(*t["changed_detected"]), *t["changed_detected"]),
            ("values_flagged_section_level", t["flagged"], "", ""),
            ("values_flagged_whole_document", t["whole_document"], "", ""),
            ("review_reduction_pct", _pct(t["whole_document"] - t["flagged"], t["whole_document"]),
             t["whole_document"] - t["flagged"], t["whole_document"]),
            ("changed_passage_recall_pct", _pct(t["changed_flagged"], t["changed_passages"]),
             t["changed_flagged"], t["changed_passages"]),
            ("flag_precision_pct", _pct(t["changed_flagged"], t["flagged"]), t["changed_flagged"], t["flagged"]),
            ("approvals_kept_pct", _pct(*t["approvals_kept"]), *t["approvals_kept"]),
            ("cases_without_decision_nodes_in_impact_pct", _pct(*t["no_decision_nodes"]), *t["no_decision_nodes"]),
            ("readiness_pending_iff_flagged_pct", _pct(*t["readiness_pending_iff_flagged"]),
             *t["readiness_pending_iff_flagged"]),
            ("selective_equals_full_pct", _pct(*t["verified"]), *t["verified"])]
    return rows


def write_tables(result: ProvenanceResult, results_dir: Path) -> dict[str, bool]:
    out = Path(results_dir) / "bench" / result.suite
    cases = tsv(["case", "document", "family", "flagged", "expected", "agree", "changed_passages", "changed_detected",
                 "whole_document", "reanchored", "sections_differing", "readiness", "approval_kept",
                 "no_decision_nodes", "verified"],
                [[r["case"], r["document"], r["family"], ",".join(r["flagged"]), ",".join(r["expected"]), r["agree"],
                  ",".join(r["expected_changed"]), ",".join(r["changed_detected"]), r["whole_document"],
                  r["reanchored"], r["sections_differing"], r["readiness"], r["approval_kept"], r["no_decision_nodes"],
                  r["verified"]] for r in result.rows])
    summary = tsv(["metric", "value", "numerator", "denominator"], summary_rows(result))
    return {str(out / "cases.tsv"): write_if_changed(out / "cases.tsv", cases),
            str(out / "summary.tsv"): write_if_changed(out / "summary.tsv", summary)}


def summary_text(result: ProvenanceResult) -> str:
    lines = [f"provenance suite {result.suite} at {result.parent}: {len(result.rows)} cases"]
    for name, value, num, den in summary_rows(result)[1:]:
        lines.append(f"  {name}: {value}" + (f" ({num}/{den})" if den != "" else ""))
    for r in result.rows:
        if not r["agree"]:
            lines.append(f"  DISAGREE {r['case']}: flagged {r['flagged']}, labelled {r['expected']}")
        if not r["changed_agree"]:
            lines.append(f"  DISAGREE {r['case']}: changed passages {r['changed_detected']}, labelled "
                         f"{r['expected_changed']}")
    return "\n".join(lines) + "\n"


__all__ = ["ProvenanceResult", "apply_edits", "run_provenance", "summary_rows", "summary_text", "write_docx",
           "write_tables", "normalize_text"]
