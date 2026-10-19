"""Index reports (Word, PDF, Markdown, text) into sections with content digests, and locate quoted values.

A document is a sequence of sections: the body text between one heading and the next (of any level), with
its heading path (e.g. "3 Results > 3.2 Vehicle B"), the page of each paragraph (PDF) and a digest of its
normalized body text. The heading is not part of a section's digest, so renaming or renumbering a section
does not make it a different section, while editing one of its sentences does. The document digest covers
every heading and section, so any edit makes a new version.

A locator pins a value to the passage it came from: the document and section digests, the section's id
and heading path, the page and paragraph, the exact quote with a little text either side (a W3C Web
Annotation TextQuoteSelector) and its character span in the section text (a TextPositionSelector).
`locate` returns a locator only for a quote that is in the document, so a fabricated quote is refused, and
`reanchor` moves a locator to a new version of the document only when its section's text is unchanged.

Text is compared after normalization (superscripts kept as exponents, NFKC, typographic quotes and dashes
folded, whitespace collapsed, PDF line-break hyphenation joined), so a quote copied from `pdo docs show`
always matches. Matching and digests also ignore a hyphen between two letters (the second lowercase), so a
word hyphenated across a PDF line ("mixed-" + "terrain", or "reliabil-" + "ity") matches its source. A quote
must start and end on whole words and numbers: "0.9" is not a quote of "0.98".

Section ids are the numbers printed in the headings ("3.2 Vehicle B" has id "3.2") when the headings carry
numbers; unnumbered headings in such a document get "s" and their outline position ("s1"), and a document
without printed numbers uses outline positions ("1", "1.2").
"""

from __future__ import annotations

import difflib
import re
import statistics
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from pdo.canonical import digest

PARSER_VERSION = 2
CONTEXT = 32  # characters of prefix and suffix kept around a quote
FORMATS = {".docx": "docx", ".pdf": "pdf", ".md": "markdown", ".markdown": "markdown", ".txt": "text"}

_FOLD = str.maketrans({
    "‘": "'", "’": "'", "‚": "'", "‛": "'", "′": "'",
    "“": '"', "”": '"', "„": '"', "″": '"',
    "‐": "-", "‑": "-", "‒": "-", "–": "-", "—": "-", "−": "-",
    " ": " ", " ": " ", " ": " ", " ": " ",
    "­": None, "​": None, "﻿": None,
})
_NUMBERED = re.compile(r"^(\d+(?:\.\d+)*)\.?\s+([A-Z][^\n]{0,78})$")
_PRINTED = re.compile(r"^(\d+(?:\.\d+)*)\.?\s")
_SUPER = str.maketrans("⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺", "0123456789-+")
_SUB = str.maketrans("₀₁₂₃₄₅₆₇₈₉₋₊", "0123456789-+")


class DocumentError(ValueError):
    """The document cannot be read (missing file, unsupported format, or the docs extra is not installed)."""

    def __init__(self, message: str, hint: str = "Documents are .docx, .pdf, .md or .txt files; Word and PDF need "
                                                 "the docs extra (uv sync --extra docs)."):
        super().__init__(message)
        self.hint = hint


class QuoteError(ValueError):
    """A quote is not in the document, or is in it more than once."""

    def __init__(self, message: str, hint: str):
        super().__init__(message)
        self.hint = hint


def normalize_text(text: str) -> str:
    # Superscript and subscript digits stay distinct from the number before them (10⁵ is "10^5", not 105),
    # and a vulgar fraction from the whole number before it (1½ is "1 1/2", not 11/2).
    text = re.sub("[⁰¹²³⁴⁵⁶⁷⁸⁹⁻⁺]+", lambda m: "^" + m.group(0).translate(_SUPER), text)
    text = "".join(" " + unicodedata.normalize("NFKC", ch).replace("⁄", "/")
                   if unicodedata.decomposition(ch).startswith("<fraction>") else ch for ch in text)
    text = re.sub("[₀₁₂₃₄₅₆₇₈₉₋₊]+", lambda m: "_" + m.group(0).translate(_SUB), text)
    return " ".join(unicodedata.normalize("NFKC", text).translate(_FOLD).split())


def _key(text: str) -> tuple[str, list[int]]:
    """The text that matching and digests compare: a hyphen between two letters (the second lowercase) is
    dropped, so line-break hyphenation joined by a PDF parser equals the source. Returns it with, for each of
    its characters, the index in `text`."""
    out, where = [], []
    for i, ch in enumerate(text):
        if ch == "-" and 0 < i < len(text) - 1 and text[i - 1].isalpha() and text[i + 1].islower():
            continue
        out.append(ch)
        where.append(i)
    return "".join(out), where


def _bounded(text: str, start: int, end: int) -> bool:
    """Whether text[start:end] starts and ends on whole words and numbers (it never cuts 0.98 to 0.9)."""
    before = text[start - 1] if start > 0 else " "
    after = text[end] if end < len(text) else " "
    first, last = text[start], text[end - 1]
    if (first.isalnum() and before.isalnum()) or (last.isalnum() and after.isalnum()):
        return False
    if last.isdigit() and after in ".," and text[end + 1:end + 2].isdigit():
        return False
    if last.isdigit() and (after == "^" or re.match(r"\s?[×x·*]\s?10", text[end:end + 5])
                           or re.match(r"[eE][-+]?\d", text[end:end + 3])):
        return False  # 1.2 of 1.2×10^-3, 10 of 10^5, 1.2 of 1.2e-3
    if first.isdigit() and before in "^eE" and text[start - 2:start - 1].isdigit():
        return False
    if last in ".," and text[end - 2:end - 1].isdigit() and after.isdigit():
        return False
    if first.isdigit() and (before == "-" or (before in ".," and text[start - 2:start - 1].isdigit())):
        return False
    return not (first in ".," and text[start + 1:start + 2].isdigit() and before.isdigit())


@dataclass
class Section:
    id: str  # printed number ("3.2") or outline position; "0" holds any text before the first heading
    heading: str
    level: int
    path: list[str]
    paragraphs: list[str]
    pages: list[int | None]  # page of each paragraph (PDF), else None
    digest: str = ""

    def __post_init__(self) -> None:
        # Paragraph breaks are not part of the digest: PDF paragraphs are reconstructed from line lengths, so
        # only the words decide whether two versions of a section are the same.
        self.digest = self.digest or digest(_key(" ".join(self.paragraphs))[0])

    @property
    def text(self) -> str:
        return "\n".join(self.paragraphs)

    @property
    def path_text(self) -> str:
        return " > ".join(self.path)

    @property
    def page_range(self) -> tuple[int | None, int | None]:
        pages = [p for p in self.pages if p is not None]
        return (min(pages), max(pages)) if pages else (None, None)

    def summary(self) -> dict:
        first, last = self.page_range
        return {"id": self.id, "heading": self.heading, "level": self.level, "path": self.path_text,
                "pages": [first, last] if first is not None else None, "paragraphs": len(self.paragraphs),
                "characters": len(self.text), "digest": self.digest}

    def to_dict(self) -> dict:
        return {**self.summary(), "path": self.path, "text": self.paragraphs, "page_of": self.pages}

    @classmethod
    def from_dict(cls, data: dict) -> Section:
        return cls(data["id"], data["heading"], data["level"], list(data["path"]), list(data["text"]),
                   list(data["page_of"]), data["digest"])


@dataclass
class DocumentIndex:
    path: str
    format: str
    sections: list[Section]
    title: str | None = None
    pages: int | None = None
    parser: str = f"pdo.docs {PARSER_VERSION}"
    digest: str = ""
    warnings: list[str] = field(default_factory=list)

    def __post_init__(self) -> None:
        self.digest = self.digest or digest([[s.path, s.digest] for s in self.sections])

    def section(self, ref: str) -> Section:
        """A section by id ("3.2"), heading ("Vehicle B") or heading path ("3 Results > 3.2 Vehicle B")."""
        wanted = normalize_text(ref)
        def bare(heading: str) -> str:  # "3.2 Vehicle B" -> "vehicle b"
            return re.sub(r"^\d+(?:\.\d+)*\.?\s+", "", normalize_text(heading)).lower()

        for test in (lambda s: s.id == ref, lambda s: normalize_text(s.path_text) == wanted,
                     lambda s: normalize_text(s.heading) == wanted,
                     lambda s: normalize_text(s.heading).lower() == wanted.lower(),
                     lambda s: bare(s.heading) == bare(wanted)):
            found = [s for s in self.sections if test(s)]
            if len(found) == 1:
                return found[0]
            if len(found) > 1:
                raise QuoteError(f"{ref!r} names {len(found)} sections of {self.path}",
                                 "Use the section id: " + ", ".join(f"{s.id} ({s.path_text})" for s in found))
        raise QuoteError(f"no section {ref!r} in {self.path}",
                         "Sections: " + "; ".join(f"{s.id} {s.heading}" for s in self.sections))

    def summary(self) -> dict:
        return {"path": self.path, "format": self.format, "title": self.title, "pages": self.pages,
                "digest": self.digest, "parser": self.parser, "sections": [s.summary() for s in self.sections],
                "warnings": self.warnings}

    def to_dict(self) -> dict:
        return {**self.summary(), "sections": [s.to_dict() for s in self.sections]}

    @classmethod
    def from_dict(cls, data: dict) -> DocumentIndex:
        return cls(data["path"], data["format"], [Section.from_dict(s) for s in data["sections"]], data.get("title"),
                   data.get("pages"), data.get("parser", ""), data["digest"], list(data.get("warnings", [])))


# --- building sections ------------------------------------------------------------------------------

class _Builder:
    """Collects (heading level, heading text) and paragraph events into sections with outline ids."""

    def __init__(self) -> None:
        self.sections: list[Section] = []
        self.stack: list[tuple[int, str]] = []  # (heading level, heading text) of the open ancestors
        self.counters: list[int] = []
        self.current = Section("0", "(before the first heading)", 0, [], [], [])

    def heading(self, level: int, text: str) -> None:
        text = normalize_text(text)
        if not text:  # an empty heading paragraph is not a section
            return
        self._close()
        while self.stack and self.stack[-1][0] >= level:
            self.stack.pop()
        depth = len(self.stack) + 1
        self.counters = (self.counters + [0] * depth)[:depth]
        self.counters[-1] += 1
        self.stack.append((level, text))
        self.current = Section(".".join(map(str, self.counters)), text, depth, [h for _, h in self.stack], [], [])

    def paragraph(self, text: str, page: int | None = None) -> None:
        text = normalize_text(text)
        if text:
            self.current.paragraphs.append(text)
            self.current.pages.append(page)

    def _close(self) -> None:
        s = self.current
        if s.id != "0" or s.paragraphs:
            self.sections.append(Section(s.id, s.heading, s.level, s.path, s.paragraphs, s.pages))

    def finish(self) -> list[Section]:
        self._close()
        numbers = [(m.group(1) if (m := _PRINTED.match(s.heading)) and s.id != "0" else None) for s in self.sections]
        if any(numbers):  # ids are the printed numbers; unnumbered or repeated ones get "s" + outline position
            counts = Counter(n for n in numbers if n)
            counts.update(s.id for s in self.sections if s.id == "0")
            for s, n in zip(self.sections, numbers, strict=True):
                if s.id != "0":
                    s.id = n if n and counts[n] == 1 else "s" + s.id
        return self.sections


def index_document(path: str | Path) -> DocumentIndex:
    """Parse a report into sections. Word headings come from heading styles; PDF headings from the outline
    (bookmarks), else from numbered heading lines; Markdown from # headings."""
    p = Path(path)
    if not p.is_file():
        raise DocumentError(f"no document at {p}")
    fmt = FORMATS.get(p.suffix.lower())
    if fmt is None:
        raise DocumentError(f"unsupported document type {p.suffix!r}: use .docx, .pdf, .md or .txt")
    return {"docx": _docx, "pdf": _pdf, "markdown": _markdown, "text": _text}[fmt](p)


def _docx(path: Path) -> DocumentIndex:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - extra not installed
        raise DocumentError("reading Word documents needs python-docx: uv sync --extra docs") from exc
    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises many types for a bad package
        raise DocumentError(f"{path} is not a readable Word document: {exc}") from exc
    b = _Builder()
    title = normalize_text(document.core_properties.title or "") or None
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            level = _heading_level(para.style)
            if level == 0:
                title = title or normalize_text(para.text) or None
            elif level:
                b.heading(level, para.text)
            else:
                b.paragraph(para.text)
        elif tag == "tbl":
            for row in Table(child, document).rows:
                seen, cells = set(), []
                for cell in row.cells:  # a merged cell appears once per grid column it spans
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    text = normalize_text(cell.text)
                    if text:
                        cells.append(text)
                b.paragraph(" | ".join(cells))
    return DocumentIndex(str(path), "docx", b.finish(), title=title, parser=f"pdo.docs {PARSER_VERSION} (python-docx)")


def _heading_level(style) -> int | None:
    """0 for a title, n for "Heading n" (or a style based on one), None for body text."""
    while style is not None:
        name = (style.name or "").strip()
        if name == "Title":
            return 0
        m = re.fullmatch(r"(?i)heading\s*(\d)", name)
        if m:
            return int(m.group(1))
        style = style.base_style
    return None


def _pdf(path: Path) -> DocumentIndex:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover - extra not installed
        raise DocumentError("reading PDF documents needs pypdf: uv sync --extra docs") from exc
    try:
        reader = PdfReader(str(path))
        pages = [_page_lines(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - pypdf raises many types for a damaged file
        raise DocumentError(f"{path} is not a readable PDF: {exc}") from exc
    warnings: list[str] = []
    title = normalize_text((reader.metadata.title if reader.metadata else None) or "") or None
    outline = _outline(reader)
    pages = _strip_furniture(pages, title, [entry for _, entry, _ in outline])
    headings: dict[tuple[int, int], tuple[int, str]] = {}  # (page, line) -> (level, heading)
    if outline:
        for level, entry, page in outline:
            where = _find_heading_line(pages, entry, page, headings)
            if where is None:
                warnings.append(f"outline entry {entry!r} (page {page + 1}) not found in the page text")
            else:
                headings[where] = (level, pages[where[0]][where[1]])
    else:
        lines = [((pi, li), line) for pi, page_lines in enumerate(pages) for li, line in enumerate(page_lines)]
        headings = _numbered_headings(lines)
        warnings.append("no outline (bookmarks): headings are numbered lines whose numbers continue the outline; "
                        "export the PDF with bookmarks for a reliable structure" if headings else
                        "no outline and no numbered headings: the document is one section")
    b = _Builder()
    para: list[str] = []  # a paragraph may continue onto the next page; it is attributed to its first page
    start = [1]

    def flush() -> None:
        if para:
            b.paragraph(_dehyphenate(para), start[0])
            para.clear()

    for pi, lines in enumerate(pages):
        width = statistics.median([len(x) for x in lines]) if lines else 0
        for li, line in enumerate(lines):
            if (pi, li) in headings:
                flush()
                b.heading(*headings[(pi, li)])
                continue
            if not para:
                start[0] = pi + 1
            para.append(line)
            ends_sentence = line.rstrip().endswith((".", ":", ";", "!", "?"))
            if (ends_sentence and len(line) < 0.85 * width) or len(line) < 0.5 * width:
                flush()
    flush()
    return DocumentIndex(str(path), "pdf", b.finish(), title=title, pages=len(pages),
                         parser=f"pdo.docs {PARSER_VERSION} (pypdf)", warnings=warnings)


def _page_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def _strip_furniture(pages: list[list[str]], title: str | None, headings: list[str] = ()) -> list[list[str]]:
    """Drop page numbers, running headers and footers, and the printed title, which would otherwise land in
    the middle of whichever section spans a page break. A line is furniture only when it is the first or last
    line of most pages (both of two) and identical on each, or differs only in a number that tracks the page
    (a folio, "Page 3 of 9"); a line that prints an outline heading never is."""
    n_pages = len(pages)
    titles = {normalize_text(h).lower() for h in headings}

    def norm(line: str) -> str:
        return normalize_text(line).lower()

    def tracks_page(entries: list[tuple[int, str]]) -> bool:
        """Every entry's numbers include its page number plus one fixed offset (or the page count)."""
        offsets = []
        for pi, line in entries:  # standalone integers only: never a part of 0.97 or 2,000
            nums = [int(x) for x in re.findall(r"(?<![\d.,])\d+(?![.,]?\d)", line)]
            if not nums:
                return False
            offsets.append({x - (pi + 1) for x in nums} | ({"total"} if n_pages in nums else set()))
        common = set.intersection(*offsets) - {"total"}
        return bool(common)

    furniture: set[tuple[int, int]] = set()  # (page, position)
    for pos in (0, -1):
        groups: dict[str, list[tuple[int, str]]] = {}
        for pi, lines in enumerate(pages):
            if lines and norm(lines[pos]) not in titles:
                groups.setdefault(re.sub(r"\d+", "#", norm(lines[pos])), []).append((pi, lines[pos]))
        for entries in groups.values():
            identical = len({norm(line) for _, line in entries}) == 1
            numbered = "#" in re.sub(r"\d+", "#", norm(entries[0][1]))
            folio_only = re.fullmatch(r"(?i)(page\s+)?\d{1,4}(\s+(of|/)\s+\d{1,4})?", entries[0][1].strip())
            if len(entries) >= 2 and len(entries) > n_pages / 2 and (
                    identical or (numbered and tracks_page(entries))):
                furniture |= {(pi, pos) for pi, _ in entries}
            elif folio_only and all(int(re.findall(r"\d+", line)[0]) == pi + 1 for pi, line in entries):
                furniture |= {(pi, pos) for pi, _ in entries}  # a page number on some pages only
    out = []
    for pi, lines in enumerate(pages):
        kept = list(lines)
        if (pi, -1) in furniture and kept:
            kept.pop()
        if (pi, 0) in furniture and kept:
            kept.pop(0)
        if pi == 0 and title and kept and normalize_text(kept[0]) == title:
            kept.pop(0)
        out.append(kept)
    return out


def _dehyphenate(lines: list[str]) -> str:
    """Join a paragraph's lines. A word broken at the line end ("reliabil-" + "ity") is joined without its
    hyphen; a hyphen that belongs to the text ("MIL-" + "STD-810", "2,000-" + "mile") is kept, with no space."""
    out = lines[0]
    for line in lines[1:]:
        if re.search(r"[A-Za-z]-$", out) and line[:1].islower():
            out = out[:-1] + line
        elif re.search(r"[A-Za-z0-9]-$", out) and line[:1].isalnum():
            out = out + line
        else:
            out = out + " " + line
    return out


def _outline(reader) -> list[tuple[int, str, int]]:
    """(level, title, page index) for every bookmark, depth first."""
    out: list[tuple[int, str, int]] = []

    def walk(entries, level: int) -> None:
        for entry in entries:
            if isinstance(entry, list):
                walk(entry, level + 1)
                continue
            try:
                page = reader.get_destination_page_number(entry)
            except Exception:  # noqa: BLE001 - a bookmark without a page is skipped
                continue
            out.append((level, normalize_text(str(entry.title)), page))

    try:
        walk(reader.outline, 1)
    except Exception:  # noqa: BLE001 - damaged outline: fall back to numbered headings
        return []
    return out


def _numbered_headings(lines: list[tuple[Any, str]]) -> dict[Any, tuple[int, str]]:
    """Heading lines of a document without an outline: numbered lines ("3.2 Results") that form the longest
    chain in which each number continues the previous one (next sibling, first child, or a later ancestor's
    sibling), so a numbered table row or list item does not become a heading."""
    cands = []
    for where, line in lines:
        m = _NUMBERED.match(line)
        if m and not line.rstrip().endswith((".", ",", ";", ":")):
            cands.append((where, tuple(int(x) for x in m.group(1).split(".")), line))

    def step(prev: tuple | None, cur: tuple) -> float:
        """1 for the next number, 0.6 for a later one (a removed section), 0 if it cannot follow."""
        if prev is None:
            return 1.0 if cur[-1] in (0, 1) else 0.6 if len(cur) == 1 else 0.0
        if cur == prev + (1,) or any(cur == prev[:k] + (prev[k] + 1,) for k in range(len(prev))):
            return 1.0
        if any(len(cur) == k + 1 and cur[:k] == prev[:k] and cur[k] > prev[k] for k in range(len(prev))):
            return 0.6
        return 0.0

    def cost(line: str) -> float:  # a long line is more likely a wrapped sentence than a heading
        return len(line) / 1000

    best, back = [0.0] * len(cands), [-1] * len(cands)
    for i, (_, num, line) in enumerate(cands):
        start = step(None, num)
        if start:
            best[i] = start - cost(line)
        for j in range(i):
            gain = step(cands[j][1], num) if best[j] else 0.0
            if gain and best[j] + gain - cost(line) > best[i]:
                best[i], back[i] = best[j] + gain - cost(line), j
    if not any(best):
        return {}
    i = max(range(len(cands)), key=lambda k: (best[k], k))
    chain = []
    while i >= 0:
        chain.append(cands[i])
        i = back[i]
    return {where: (len(num), line) for where, num, line in chain}


def _find_heading_line(pages: list[list[str]], title: str, page: int, taken) -> tuple[int, int] | None:
    """The line that prints a bookmark's title: on its page (or the next), equal to the title, optionally
    after a section number."""
    want = normalize_text(title).lower()
    for pi in (page, page + 1):
        if pi >= len(pages):
            break
        for li, line in enumerate(pages[pi]):
            if (pi, li) in taken:
                continue
            text = normalize_text(line).lower()
            if text == want or re.fullmatch(r"\d+(?:\.\d+)*\.?\s+" + re.escape(want), text) \
                    or re.fullmatch(r"\d+(?:\.\d+)*\.?\s+" + re.escape(text), want):
                return (pi, li)
    return None


def _inline(text: str) -> str:
    """Markdown inline markup to the text a renderer shows (links and images keep their text)."""
    text = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\[[^\]]*\]", r"\1", text)
    text = re.sub(r"`([^`]*)`", r"\1", text)
    text = re.sub(r"(\*\*|__)(?=\S)(.+?)(?<=\S)\1", r"\2", text)
    text = re.sub(r"(?<![\w*])\*(?=\S)(.+?)(?<=\S)\*(?![\w*])", r"\1", text)
    text = re.sub(r"(?<![\w_])_(?=\S)(.+?)(?<=\S)_(?![\w_])", r"\1", text)
    text = re.sub(r"</?[A-Za-z][^>\n]*>", "", text)
    return re.sub(r"\\([\\`*_{}\[\]()#+\-.!|])", r"\1", text)


def _markdown(path: Path) -> DocumentIndex:
    text = path.read_text(encoding="utf-8")
    if text.startswith("---\n"):  # front matter
        end = text.find("\n---", 4)
        text = text[end + 4:] if end > 0 else text
    text = re.sub(r"<!--.*?-->", "", text, flags=re.S)
    b = _Builder()
    para: list[str] = []
    fence = ""

    def flush() -> None:
        if para:
            b.paragraph(_inline(" ".join(para)))
            para.clear()

    for raw in text.splitlines():
        line = raw.strip()
        marker = re.match(r"^(```|~~~)", line)
        if marker and (not fence or line.startswith(fence)):
            flush()
            fence = "" if fence else marker.group(1)
            continue
        if fence:
            if line:
                para.append(line)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*?)\s*#*$", line)
        setext = re.fullmatch(r"=+|-+", line)
        if heading:
            flush()
            b.heading(len(heading.group(1)), _inline(heading.group(2)))
        elif setext and len(para) == 1:  # "Title" over "=====" or "-----"
            title = para.pop()
            b.heading(1 if line[0] == "=" else 2, _inline(title))
        elif setext:
            flush()  # a thematic break
        elif not line:
            flush()
        elif line.startswith("|"):
            flush()
            cells = [_inline(c.strip()) for c in line.strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c):
                b.paragraph(" | ".join(c for c in cells if c))
        elif re.match(r"^([-*+]|\d+[.)])\s+", line):
            flush()
            para.append(re.sub(r"^([-*+]|\d+[.)])\s+", "", line))
        else:
            para.append(line)
    flush()
    return DocumentIndex(str(path), "markdown", b.finish())


def _text(path: Path) -> DocumentIndex:
    b = _Builder()
    para: list[str] = []
    raw_lines = [raw.strip() for raw in path.read_text(encoding="utf-8").splitlines()] + [""]
    headings = _numbered_headings([(i, line) for i, line in enumerate(raw_lines) if line])
    for i, line in enumerate(raw_lines):
        if i in headings:
            if para:
                b.paragraph(" ".join(para))
                para.clear()
            b.heading(*headings[i])
        elif line:
            para.append(line)
        elif para:
            b.paragraph(" ".join(para))
            para.clear()
    return DocumentIndex(str(path), "text", b.finish())


# --- locating quotes ------------------------------------------------------------------------------

def _occurrences(section: Section, quote: str) -> list[tuple[int, int]]:
    """(start, end) spans of a normalized quote in the section text, compared hyphen-insensitively and kept
    only where the span starts and ends on whole words and numbers."""
    flat = section.text.replace("\n", " ")
    key, where = _key(flat)
    qkey = _key(quote)[0]
    out = []
    i = key.find(qkey) if qkey else -1
    while i >= 0:
        start, end = where[i], where[i + len(qkey) - 1] + 1
        if _bounded(flat, start, end):
            out.append((start, end))
        i = key.find(qkey, i + 1)
    return out


def _locator(index: DocumentIndex, section: Section, start: int, end: int) -> dict:
    text = section.text
    paragraph = text.count("\n", 0, start)
    return {
        "document_digest": index.digest, "section": section.id, "section_path": section.path_text,
        "section_digest": section.digest, "page": section.pages[paragraph] if section.pages else None,
        "paragraph": paragraph + 1, "quote": text[start:end].replace("\n", " "),
        "quote_prefix": text[max(0, start - CONTEXT):start].replace("\n", " "),
        "quote_suffix": text[end:end + CONTEXT].replace("\n", " "),
        "char_start": start, "char_end": end,
    }


def locate(index: DocumentIndex, quote: str, section: str | None = None) -> dict:
    """A locator for `quote`, which must occur exactly once in the document (or in `section`), starting and
    ending on whole words and numbers."""
    q = normalize_text(quote)
    if len(q) < 4:
        raise QuoteError("the quote is too short to locate a value", "Quote the sentence or table row that states it.")
    scope = [index.section(section)] if section else index.sections
    hits = [(s, start, end) for s in scope for start, end in _occurrences(s, q)]
    if len(hits) == 1:
        return _locator(index, *hits[0])
    where = f"section {scope[0].id}" if section else index.path
    if hits:
        places = sorted({f"{s.id} ({s.heading})" for s, _, _ in hits})
        raise QuoteError(f"the quote occurs {len(hits)} times in {where}",
                         "Quote more of the passage" + ("" if section else ", or name the section") + ": "
                         + ", ".join(places))
    raise QuoteError(f"the quote is not in {where}, as whole words and numbers: {q!r}", _closest(scope, q))


def _closest(sections: list[Section], quote: str) -> str:
    best: tuple[float, Section | None, str] = (0.0, None, "")
    for s in sections:
        flat = s.text.replace("\n", " ")
        if not flat:
            continue
        m = difflib.SequenceMatcher(None, quote, flat, autojunk=False).find_longest_match(0, len(quote), 0, len(flat))
        score = m.size / len(quote)
        if score > best[0]:
            lo = max(0, m.b - m.a - 10)
            best = (score, s, flat[lo:lo + len(quote) + 20])
    if best[1] is None or best[0] < 0.3:
        return "Read the section text with pdo docs show and copy the passage exactly."
    return (f"The closest passage is in section {best[1].id} ({best[1].heading}): {best[2]!r}. Copy the passage "
            "exactly as pdo docs show prints it, including whole numbers; never paraphrase a quote.")


def verify_locator(index: DocumentIndex, locator: dict) -> list[str]:
    """Problems with a locator against this document (empty if it holds): another version, a missing or
    changed section, a quote that is not at its recorded span, or any recorded field (heading path, page,
    paragraph, prefix, suffix) that the document does not give."""
    problems = []
    if locator.get("document_digest") != index.digest:
        problems.append("made against another version of the document")
    try:
        section = index.section(locator["section"])
    except (QuoteError, KeyError):
        return problems + [f"no section {locator.get('section')} in this version"]
    if locator.get("section_digest") != section.digest:
        problems.append(f"section {section.id} has changed")
    spans = _occurrences(section, normalize_text(locator.get("quote") or ""))
    span = (locator.get("char_start"), locator.get("char_end"))
    if not spans:
        return problems + [f"the quote is not in section {section.id}"]
    if span not in spans:
        return problems + [f"the quote is in section {section.id} but not at characters {span[0]}-{span[1]}"]
    expected = _locator(index, section, *span)
    for k in ("section_path", "page", "paragraph", "quote", "quote_prefix", "quote_suffix"):
        if (locator.get(k) or None) != (expected[k] or None):
            problems.append(f"{k} is {locator.get(k)!r}, but the document gives {expected[k]!r}")
    return problems


def _bare(heading: str) -> str:
    return _PRINTED.sub("", normalize_text(heading), count=1).lower()


def _number(heading: str) -> str | None:
    m = _PRINTED.match(normalize_text(heading))
    return m.group(1) if m else None


def _successor(new: DocumentIndex, path: str | None, digest_: str | None, old_paths: set[str]) -> tuple:
    """Where an old section is in a new version: (section, status). It is found by identity, never by its
    text alone: the section at the same heading path; else the one with the same heading, its number aside
    (renumbered: moved or changed); else the one with the same printed number and the same text (renamed:
    moved). So a changed section never follows its old text into another section (an appendix keeping
    superseded text, say), and a renamed section must keep its text to be recognised. When nothing
    identifies it, it is ambiguous if its old text appears elsewhere, changed if a section with its number
    remains, else removed."""
    heading = (path or "").split(" > ")[-1]
    same_path = [s for s in new.sections if s.path_text == path]
    if same_path:
        hit = next((s for s in same_path if s.digest == digest_), None)
        return (hit, "unchanged") if hit else (same_path[0], "changed")
    others = [s for s in new.sections if s.path_text not in old_paths]  # not the continuation of an old section
    renumbered = [s for s in others if _bare(s.heading) == _bare(heading)]
    if len(renumbered) == 1:
        return renumbered[0], "moved" if renumbered[0].digest == digest_ else "changed"
    renamed = [s for s in others if _number(heading) and _number(s.heading) == _number(heading)]
    if len(renamed) == 1 and renamed[0].digest == digest_:
        return renamed[0], "moved"
    if any(s.digest == digest_ and s.paragraphs for s in others):
        return None, "ambiguous"  # its old text is elsewhere, but nothing identifies that section as this one
    if len(renamed) == 1:
        return renamed[0], "changed"  # renamed and edited
    return None, "removed"


def reanchor(locator: dict, index: DocumentIndex, old: DocumentIndex | None = None) -> tuple[dict | None, dict]:
    """Move a locator to a new version of its document, only when its section's text is unchanged there (the
    heading or position may differ; see _successor). `old`, the version the locator was made against, makes
    moves more precise. Returns (new locator or None, what happened)."""
    old_paths = {s.path_text for s in old.sections} if old is not None else set()
    section, status = _successor(index, locator.get("section_path"), locator.get("section_digest"), old_paths)
    quote = normalize_text(locator.get("quote") or "")
    if section is not None and status in ("unchanged", "moved"):
        spans = _occurrences(section, quote)
        if spans:
            span = (locator.get("char_start"), locator.get("char_end"))
            return _locator(index, section, *(span if span in spans else spans[0])), \
                {"status": status, "section": section.id}
    found = sorted({s.id for s in index.sections if quote and _occurrences(s, quote)})
    successor = section.id if section is not None else None
    return None, {"status": status if status not in ("unchanged", "moved") else "changed", "section": successor,
                  "quote_found_in": found, "quote_in_section": successor in found}


def compare(old: DocumentIndex, new: DocumentIndex) -> list[dict]:
    """How each section of `old` fared in `new` (by the rules of _successor, matched one to one): unchanged,
    changed, moved, ambiguous or removed; plus the sections added in `new`."""
    old_paths = {s.path_text for s in old.sections}
    out, matched = [], set()
    for s in old.sections:
        if not s.paragraphs:
            continue
        candidates = DocumentIndex(new.path, new.format, [t for t in new.sections if t.id not in matched], digest="-")
        section, status = _successor(candidates, s.path_text, s.digest, old_paths)
        if section is not None:
            matched.add(section.id)
        out.append({"section": s.id, "path": s.path_text, "status": status,
                    "now": section.id if section is not None else None})
    for t in new.sections:
        if t.paragraphs and t.id not in matched:
            out.append({"section": None, "path": t.path_text, "status": "added", "now": t.id})
    return out
